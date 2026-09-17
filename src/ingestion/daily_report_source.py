"""Источник суточного рапорта по добыче (CSV/Excel), скважина x сутки.

Отдельно от csv_source.py (измерения tag=value) и monthly_production_source.py
(месячные агрегаты): здесь готовая суточная запись с дебитом и типом
источника замера (rate/accumulated), которая ещё требует расчёта добычи —
сам расчёт в src/calc/daily_report.py, здесь только разбор файла.

Маппинг колонок задаётся YAML (src/ingestion/mappings/daily_report/), т.к.
формат суточного рапорта у каждой компании свой.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterator

import pandas as pd
import yaml
from pydantic import BaseModel

from src.ingestion.base import SourceValidationError

# Обязательные поля строки — без них расчёт в принципе невозможен.
# source_type НЕ входит сюда специально: его отсутствие — не ошибка парсинга,
# а состояние "требует уточнения", которое считает src/calc/daily_report.py.
REQUIRED_FIELDS = ("hours_on", "q_liquid_m3", "water_cut_pct")
OPTIONAL_FIELDS = ("q_gas_thousand_m3",)


class DailyReportMapping(BaseModel):
    external_system: str
    well_id_aliases: list[str]
    date_column_aliases: list[str]
    source_type_column_aliases: list[str] = []
    comment_column_aliases: list[str] = []
    equipment_type_column_aliases: list[str] = []
    header_rows: int = 1
    fields: dict[str, list[str]]

    @classmethod
    def from_yaml(cls, path: str | Path) -> "DailyReportMapping":
        with open(path, encoding="utf-8") as f:
            raw = yaml.safe_load(f)
        return cls.model_validate(raw)


class RawDailyRecord(BaseModel):
    external_well_id: str | None = None
    date: object = None  # dt.date | None
    source_type: str | None = None
    equipment_type: str | None = None
    comment: str | None = None
    values: dict[str, float] = {}
    raw: dict = {}
    parse_error: str | None = None


def _normalize_header(s: object) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _is_empty(value: object) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    return str(value).strip() == ""


def _parse_date(value: object):
    if _is_empty(value):
        return None
    text = str(value).strip()
    iso = pd.to_datetime(text, format="%Y-%m-%d", errors="coerce")
    if not pd.isna(iso):
        return iso
    parsed = pd.to_datetime(text, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed


def _parse_float(value: object) -> float | None:
    if _is_empty(value):
        return None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


class DailyReportSource:
    """CSV/Excel-источник суточного рапорта. Контракт как у
    MonthlyProductionSource (validate_schema/read/get_metadata), тоже не
    наследует общий DataSource — форма записи другая."""

    def __init__(self, path: str | Path, mapping_path: str | Path):
        self.path = Path(path)
        self.mapping = DailyReportMapping.from_yaml(mapping_path)
        self._df_cache: pd.DataFrame | None = None

    def get_metadata(self) -> dict:
        return {
            "source": "daily_csv",
            "external_system": self.mapping.external_system,
            "description": f"Суточный рапорт: {self.path.name}",
            "path": str(self.path),
        }

    def validate_schema(self) -> None:
        if not self.path.exists():
            raise SourceValidationError(f"файл не найден: {self.path}")
        try:
            df = self._load_dataframe()
        except Exception as exc:
            raise SourceValidationError(f"не удалось прочитать файл: {exc}") from exc

        header = list(df.columns)
        if self._find(header, self.mapping.well_id_aliases) is None:
            raise SourceValidationError(
                f"не найдена колонка со скважиной (ожидались: {self.mapping.well_id_aliases})"
            )
        if self._find(header, self.mapping.date_column_aliases) is None:
            raise SourceValidationError(
                f"не найдена колонка с датой (ожидались: {self.mapping.date_column_aliases})"
            )
        missing_required = [
            f for f in REQUIRED_FIELDS
            if f not in self.mapping.fields or self._find(header, self.mapping.fields[f]) is None
        ]
        if missing_required:
            raise SourceValidationError(f"не найдены обязательные колонки для полей: {missing_required}")

    def read(self) -> Iterator[RawDailyRecord]:
        df = self._load_dataframe()
        header = list(df.columns)

        well_col = self._find(header, self.mapping.well_id_aliases)
        date_col = self._find(header, self.mapping.date_column_aliases)
        source_type_col = self._find(header, self.mapping.source_type_column_aliases)
        comment_col = self._find(header, self.mapping.comment_column_aliases)
        equipment_col = self._find(header, self.mapping.equipment_type_column_aliases)
        field_cols = {field: self._find(header, aliases) for field, aliases in self.mapping.fields.items()}

        for _, row in df.iterrows():
            raw_row = row.to_dict()

            well_raw = row.get(well_col) if well_col else None
            if _is_empty(well_raw):
                yield RawDailyRecord(raw=raw_row, parse_error="пустой идентификатор скважины")
                continue

            date_raw = row.get(date_col) if date_col else None
            ts = _parse_date(date_raw)
            if ts is None:
                yield RawDailyRecord(raw=raw_row, parse_error=f"не удалось разобрать дату: {date_raw!r}")
                continue

            values: dict[str, float] = {}
            parse_error = None
            for field, col_name in field_cols.items():
                if col_name is None:
                    continue
                value_raw = row.get(col_name)
                if _is_empty(value_raw):
                    continue
                value = _parse_float(value_raw)
                if value is None:
                    parse_error = f"не удалось разобрать значение {col_name}={value_raw!r}"
                    break
                values[field] = value

            if parse_error:
                yield RawDailyRecord(raw=raw_row, parse_error=parse_error)
                continue

            missing = [f for f in REQUIRED_FIELDS if f not in values]
            if missing:
                yield RawDailyRecord(raw=raw_row, parse_error=f"не заполнены обязательные поля: {missing}")
                continue

            source_type_raw = row.get(source_type_col) if source_type_col else None
            comment_raw = row.get(comment_col) if comment_col else None
            equipment_raw = row.get(equipment_col) if equipment_col else None

            yield RawDailyRecord(
                external_well_id=str(well_raw).strip(),
                date=ts.date(),
                source_type=str(source_type_raw).strip() if not _is_empty(source_type_raw) else None,
                equipment_type=str(equipment_raw).strip() if not _is_empty(equipment_raw) else None,
                comment=str(comment_raw).strip() if not _is_empty(comment_raw) else None,
                values=values,
                raw=raw_row,
            )

    # --- внутреннее (та же многострочная шапка Excel, что и в остальных источниках) ---

    def _load_dataframe(self) -> pd.DataFrame:
        if self._df_cache is None:
            if self.path.suffix.lower() in (".xlsx", ".xlsm"):
                self._df_cache = self._load_excel()
            else:
                self._df_cache = pd.read_csv(self.path, dtype=str, keep_default_na=False)
        return self._df_cache

    def _load_excel(self) -> pd.DataFrame:
        import openpyxl

        wb = openpyxl.load_workbook(self.path, data_only=True)
        ws = wb.active
        header_rows = self.mapping.header_rows

        grid = [
            [ws.cell(r, c).value for c in range(1, ws.max_column + 1)] for r in range(1, ws.max_row + 1)
        ]
        for rng in ws.merged_cells.ranges:
            if rng.max_row > header_rows:
                continue
            anchor = grid[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[r - 1][c - 1] = anchor

        columns = []
        for c in range(len(grid[0]) if grid else 0):
            parts = [str(grid[r][c]).strip() for r in range(header_rows) if not _is_empty(grid[r][c])]
            columns.append(" ".join(dict.fromkeys(parts)))

        data_rows = [row for row in grid[header_rows:] if any(not _is_empty(v) for v in row)]
        return pd.DataFrame(data_rows, columns=columns)

    @staticmethod
    def _find(header: list[str], aliases: list[str]) -> str | None:
        normalized = {_normalize_header(h): h for h in header}
        for alias in aliases:
            match = normalized.get(_normalize_header(alias))
            if match is not None:
                return match
        return None
