"""Excel с многострочной шапкой и объединёнными ячейками — то, что реально
встречается в промысловых рапортах."""

from pathlib import Path

import openpyxl

from src.ingestion.base import Period
from src.ingestion.csv_source import CsvSource

MAPPING = Path(__file__).resolve().parents[2] / "src" / "ingestion" / "mappings" / "example_daily_report.yaml"
PERIOD = Period.parse("2024-01")


def _build_excel_with_merged_header(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active

    # "Скважина" и "Дата" — объединены по вертикали на 2 строки шапки
    ws["A1"] = "Скважина"
    ws.merge_cells("A1:A2")
    ws["B1"] = "Дата"
    ws.merge_cells("B1:B2")

    # "Дебит" — общий заголовок сверху для двух подколонок, разворачивается
    # в "Дебит нефти, т/сут" / "Дебит жидкости, т/сут" после склейки строк шапки
    ws["C1"] = "Дебит"
    ws.merge_cells("C1:D1")
    ws["C2"] = "нефти, т/сут"
    ws["D2"] = "жидкости, т/сут"

    ws.append(["0001", "2024-01-05", 18.4, 24.6])
    ws.append(["0001", "2024-01-06", 18.1, 24.9])

    wb.save(path)


def test_excel_merged_header_and_multirow_header(tmp_path, db_session, known_well):
    path = tmp_path / "report_with_merged_header.xlsx"
    _build_excel_with_merged_header(path)

    mapping_path = tmp_path / "mapping.yaml"
    mapping_path.write_text(MAPPING.read_text(encoding="utf-8").replace("header_rows: 1", "header_rows: 2"), encoding="utf-8")

    source = CsvSource(path, mapping_path)
    source.validate_schema()  # не должно падать — шапка должна разрешиться

    records = list(source.read(PERIOD))
    tags = {r.tag for r in records}
    assert "q_oil_daily" in tags
    assert "q_liquid_daily" in tags

    oil_values = sorted(r.value for r in records if r.tag == "q_oil_daily")
    assert oil_values == [18.1, 18.4]
